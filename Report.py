from Classes.Product import Product
from FileCommands import FileCommands
from datetime import datetime
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.drawing.text import ParagraphProperties, CharacterProperties


class Report:

    @staticmethod
    def style_chart(chart, colors, categories_count):

        # --- رنگ هر سری ---
        for i, series in enumerate(chart.series):
            if i < len(colors):
                series.graphicalProperties.solidFill = colors[i]
                series.graphicalProperties.line.noFill = True

        # --- نمایش عدد بالای هر میله ---
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True
        chart.dataLabels.showLegendKey = False
        chart.dataLabels.showCatName = False
        chart.dataLabels.showSerName = False
        chart.dataLabels.numFmt = "0"

        # --- موقعیت Legend (بالای نمودار) ---
        chart.legend.position = "t"
        chart.legend.overlay = False

        # --- حذف گرید‌لاین افقی (محور Y) ---
        chart.y_axis.majorGridlines = None
        chart.y_axis.delete = True

        # --- تنظیم گپ بین میله‌ها ---
        chart.gapWidth = 150
        chart.overlap = -10

        # --- سایز نمودار ---
        chart.width = 22
        chart.height = 12

        # --- استایل عنوان ---
        chart.title.tx.rich.p[0].pPr = ParagraphProperties(
            defRPr=CharacterProperties(sz=1800, b=True, solidFill="595959")
        )

        return chart

    @staticmethod
    def generate_top_customers_report():

        customers = FileCommands.get_file_data("./Database/Customers.json")

        monthly_totals = {}
        customer_totals = {}

        for customer in customers:
            name = customer["firstname"] + " " + customer["lastname"]
            for purchase in customer["bought"]:
                date = datetime.fromisoformat(purchase["date"])
                month_key = date.strftime("%Y-%m")
                month_label = date.strftime("%B").upper()

                product = Product.get_product_byID(purchase["product_id"])
                if not product:
                    continue

                price = product["price"]
                monthly_totals.setdefault(month_key, {"label": month_label, "data": {}})
                monthly_totals[month_key]["data"][name] = (
                    monthly_totals[month_key]["data"].get(name, 0) + price
                )
                customer_totals[name] = customer_totals.get(name, 0) + price

        top_customers = sorted(customer_totals.items(), key=lambda x: x[1], reverse=True)[:3]
        top_names = [name for name, _ in top_customers]

        sorted_months = sorted(monthly_totals.keys())

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Top Customers"

        sheet.append(["Month"] + top_names)

        for month_key in sorted_months:
            month_label = monthly_totals[month_key]["label"]
            row = [month_label]
            for name in top_names:
                row.append(monthly_totals[month_key]["data"].get(name, 0))
            sheet.append(row)

        chart = BarChart()
        chart.type = "col"
        chart.grouping = "clustered"
        chart.title = "TOP 3 CUSTOMERS MONTHLY PURCHASE"
        chart.x_axis.title = "MONTH"

        n_rows = len(sorted_months)

        data = Reference(sheet, min_col=2, max_col=1 + len(top_names), min_row=1, max_row=n_rows + 1)
        categories = Reference(sheet, min_col=1, min_row=2, max_row=n_rows + 1)

        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)

        Report.style_chart(chart, ["4472C4", "C0504D", "9BBB59"], n_rows)

        sheet.add_chart(chart, "F2")
        workbook.save("./Reports/Top_Customers_Monthly_Report.xlsx")

        return True


    @staticmethod
    def generate_top_sellers_report():

        sellers = FileCommands.get_file_data("./Database/Sellers.json")

        monthly_totals = {}   # month_key -> {"label": month_label, "data": {seller_name: total}}
        seller_totals = {}    # seller_name -> total فروش کل (برای تشخیص ۳ نفر برتر)

        for seller in sellers:
            name = seller["firstname"] + " " + seller["lastname"]

            for sale in seller["sold"]:
                date = datetime.fromisoformat(sale["date"])
                month_key = date.strftime("%Y-%m")
                month_label = date.strftime("%B").upper()

                product = Product.get_product_byID(sale["product_id"])
                if not product:
                    continue

                price = product["price"]

                monthly_totals.setdefault(month_key, {"label": month_label, "data": {}})
                monthly_totals[month_key]["data"][name] = (
                    monthly_totals[month_key]["data"].get(name, 0) + price
                )

                seller_totals[name] = seller_totals.get(name, 0) + price

        # پیدا کردن ۳ فروشنده برتر (بر اساس مجموع کل فروش)
        top_sellers = sorted(seller_totals.items(), key=lambda x: x[1], reverse=True)[:3]
        top_names = [name for name, _ in top_sellers]

        # مرتب سازی ماه ها به ترتیب زمانی
        sorted_months = sorted(monthly_totals.keys())

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Top Sellers"

        sheet.append(["Month"] + top_names)

        for month_key in sorted_months:
            month_label = monthly_totals[month_key]["label"]
            row = [month_label]
            for name in top_names:
                row.append(monthly_totals[month_key]["data"].get(name, 0))
            sheet.append(row)

        chart = BarChart()
        chart.type = "col"
        chart.grouping = "clustered"
        chart.title = "TOP 3 SELLERS MONTHLY SALES"
        chart.x_axis.title = "MONTH"

        n_rows = len(sorted_months)

        data = Reference(sheet, min_col=2, max_col=1 + len(top_names), min_row=1, max_row=n_rows + 1)
        categories = Reference(sheet, min_col=1, min_row=2, max_row=n_rows + 1)

        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)

        Report.style_chart(chart, ["4472C4", "C0504D", "9BBB59"], n_rows)

        sheet.add_chart(chart, "E2")

        workbook.save("./Reports/Top_Sellers_Monthly_Report.xlsx")

        return True